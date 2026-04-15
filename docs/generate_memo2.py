"""
Generates: ML_Memo2.docx
Matches the clean black/white/grey style of Memo 1.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.15)
    section.right_margin  = Inches(1.15)

# ── Helpers ───────────────────────────────────────────────────────────────────
GREY   = RGBColor(0x88, 0x88, 0x88)
BLACK  = RGBColor(0x00, 0x00, 0x00)
LGREY  = RGBColor(0xCC, 0xCC, 0xCC)

def set_font(run, name="Garamond", size=11, bold=False, italic=False, color=BLACK):
    run.font.name  = name
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    run.font.color.rgb = color

def para(text, style="Normal", align=WD_ALIGN_PARAGRAPH.LEFT,
         space_before=0, space_after=6):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = Pt(13)
    return p

def add_text(p, text, name="Garamond", size=11, bold=False,
             italic=False, color=BLACK):
    run = p.add_run(text)
    set_font(run, name=name, size=size, bold=bold, italic=italic, color=color)
    return run

def section_heading(label, title):
    """Matches Memo 1 style: '1 . S E C T I O N' with horizontal rule."""
    doc.add_paragraph()  # small spacer
    p = para("", space_before=4, space_after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spaced = "  ".join(label.upper())
    add_text(p, spaced + "   " + "  ".join(title.upper()),
             name="Garamond", size=9, bold=False, color=GREY)
    # Horizontal rule via paragraph border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)

def sub_heading(text):
    p = para("", space_before=8, space_after=2)
    add_text(p, text, size=11, bold=True)

def body(text, space_after=6):
    p = para("", space_before=0, space_after=space_after)
    add_text(p, text, size=11)
    return p

def add_rule():
    p = para("", space_before=2, space_after=2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ── Header ────────────────────────────────────────────────────────────────────
p = para("", space_before=0, space_after=2)
spaced_label = "  ".join("PROJECT UPDATE MEMO 2")
add_text(p, spaced_label, size=8, color=GREY)

p = para("", space_before=0, space_after=4)
add_text(p, "TSA Checkpoint Passenger Flow Forecasting", size=20, bold=True)

p = para("", space_before=0, space_after=0)
add_text(p, "Arend Colle, Vivek Sriram, Tiffany T. Nguyen", size=10, italic=True, color=GREY)
add_text(p, "   April 17, 2026", size=10, italic=True, color=GREY)

add_rule()

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 1 — AWS DEPLOYMENT STATUS
# ═══════════════════════════════════════════════════════════════════════════════
section_heading("1 .", "AWS Deployment Status")

body(
    "The ML system is deployed and running on AWS using two services: Amazon S3 for artifact "
    "storage and an EC2 instance (t3.medium, 2 vCPU, 4 GB RAM) hosting the Streamlit "
    "prediction interface. The application is live at http://3.85.169.163:8501 and ready "
    "for the April 22nd demo. S3 stores the trained model file (lgbmModel.pkl) and "
    "preprocessed feature cache (modelDf.csv, 3.2M rows); the EC2 instance loads the model "
    "at startup and serves live inference in under one second per request. The architecture "
    "diagram is included in Appendix A (Figure A1)."
)

body(
    "This is a simplification of the SageMaker batch pipeline proposed at midpoint. Live "
    "EC2 inference is more appropriate for a POC demo and avoids unnecessary pipeline "
    "complexity. The main pre-demo risk is AWS Academy session timeout stopping the EC2 "
    "instance; restarting it and relaunching Streamlit takes under one minute."
)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 2 — FINAL MODEL SELECTION
# ═══════════════════════════════════════════════════════════════════════════════
section_heading("2 .", "Final Model Selection")

sub_heading("Why Linear Models Cannot Solve This Problem")

body(
    "OLS, Ridge, and Lasso regression were evaluated first. All three failed to improve "
    "on the naive historical baseline (mean throughput per airport/checkpoint/hour/day-of-week/"
    "month combination) and were ruled out. The reason is structural: throughput is governed "
    "by interactions, not additive effects. ATL is busy. Friday mornings are busy. But ATL "
    "on a Friday morning in June is disproportionately busy in a way that cannot be captured "
    "by summing an ATL coefficient, a Friday coefficient, and a morning coefficient. A linear "
    "model has no mechanism to learn joint effects, so when forced to fit 3.1 million records "
    "driven by interactions it cannot represent, it produces noisier predictions than simply "
    "looking up the historical average. Ridge and Lasso regularization did not help; the "
    "problem is underfitting from a mismatched model class, not overfitting."
)

sub_heading("Why LightGBM")

body(
    "Gradient boosted trees address the interaction problem directly. Each split encodes "
    "a conditional rule (e.g., airport=ATL AND hour=8), and a sequence of hundreds of "
    "such splits across residual-targeted trees learns the joint effects that linear "
    "models cannot represent. LightGBM is the right specific implementation for this "
    "dataset for three reasons: it handles label-encoded categoricals correctly without "
    "implying false magnitude ordering; it requires no feature standardization across "
    "the mixed-scale feature set; and its histogram-based, leaf-wise growth is "
    "computationally practical on 3.1 million training records. Early stopping on the "
    "validation set automatically selects the optimal number of boosting rounds, "
    "preventing overfitting without manual tuning."
)

sub_heading("Hyperparameter Tuning")

body(
    "A grid search was conducted over 27 combinations of learning rate [0.03, 0.05, 0.1], "
    "number of leaves [63, 127, 255], and minimum child samples [30, 50, 100], evaluated "
    "by validation RMSE. Feature fraction, bagging fraction, and bagging frequency were "
    "fixed at 0.8, 0.8, and 5; these parameters are less sensitive and including them "
    "would have substantially increased search runtime. The optimal configuration was "
    "learning rate 0.05, 127 leaves, and 50 minimum child samples. The final model was "
    "retrained on training data only with the validation set used solely for early stopping "
    "(patience of 100 rounds, cap of 10,000), converging at 1,914 rounds."
)

sub_heading("Test Set Performance")

body(
    "All metrics below are computed on the held-out test set (final 60 days of data, "
    "159,021 records), which was never seen during training or tuning. Daily Pearson r "
    "follows Monmousseau et al. (2020): predictions and actuals are summed to the daily "
    "level across all airports and checkpoints before computing correlation."
)

# ── Metrics table ─────────────────────────────────────────────────────────────
tbl = doc.add_table(rows=6, cols=4)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Metric", "Naive Baseline", "LightGBM", "Staffing Relevance"]
rows_data = [
    ["R²",               "0.7968", "0.8509",
     "Proportion of throughput variance explained"],
    ["MAE (passengers)",  "147.07", "123.85",
     "Average error in original units; used to set staffing buffers"],
    ["RMSE",             "213.53", "182.93",
     "Penalizes large single-hour misses; severe under-staffing at one checkpoint-hour has major operational consequences"],
    ["Daily Pearson r",  "N/A",    "0.9398",
     "Correlation of daily totals; identifies busy vs. quiet days for system-wide planning"],
    ["vs. Baseline",     "N/A",    "-16% MAE / -14% RMSE",
     "Consistent gains on all metrics relative to the current planning standard"],
]

# Header row
for i, h in enumerate(headers):
    cell = tbl.rows[0].cells[i]
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(h)
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.name = "Garamond"
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8E8E8')
    tcPr.append(shd)

# Data rows
for r_idx, row_data in enumerate(rows_data):
    row = tbl.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        cell = row.cells[c_idx]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(9)
        run.font.name = "Garamond"
        if c_idx == 2:
            run.font.bold = True
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 2]
            else WD_ALIGN_PARAGRAPH.LEFT
        )

col_widths = [Inches(1.2), Inches(1.1), Inches(1.1), Inches(2.9)]
for row in tbl.rows:
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]

doc.add_paragraph()

sub_heading("Tradeoffs and Feature Importance")

body(
    "LightGBM is less interpretable than a linear model, which is an acceptable tradeoff "
    "for a staffing tool where the end user needs an accurate forecast, not model weights. "
    "The 16% MAE improvement (147 to 124 passengers) compounds across dozens of checkpoints "
    "and hundreds of planning slots. Feature importance validates the model's logic: "
    "checkpoint name (52.3%), hour of day (19.2%), and airport code (15.5%) account for "
    "87% of predictive gain; these are precisely the interaction terms linear models cannot "
    "represent. Flight schedule features contribute an additional 9.4%, confirming scheduled "
    "capacity provides real signal beyond calendar and location alone."
)

# ═══════════════════════════════════════════════════════════════════════════════
# SECTION 3 — USER INTERFACE PLAN
# ═══════════════════════════════════════════════════════════════════════════════
section_heading("3 .", "User Interface Plan")

body(
    "The intended user is a TSA workforce planner making staffing decisions approximately "
    "12 months ahead. The interface is a Streamlit application deployed on EC2, organized "
    "into a sidebar for inputs and a main panel for outputs."
)

body("Inputs (sidebar):")

# Input/output table
tbl2 = doc.add_table(rows=7, cols=2)
tbl2.style = "Table Grid"
tbl2.alignment = WD_TABLE_ALIGNMENT.LEFT

io_headers = ["Input", "Description"]
io_rows = [
    ["Airport",           "Dropdown: 20 major U.S. airports (ATL, LAX, DFW, …)"],
    ["Month",             "Planning month (January – December)"],
    ["Day of week",       "Monday – Sunday"],
    ["Federal holiday",   "Checkbox: marks the date as a U.S. federal holiday"],
    ["Days to holiday",   "Slider (0–30): proximity to nearest holiday"],
    ["Flight schedule",   "Hourly departure count, avg seats/flight, avg load factor"],
]

for i, h in enumerate(io_headers):
    cell = tbl2.rows[0].cells[i]
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(h)
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.name = "Garamond"
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8E8E8')
    tcPr.append(shd)

for r_idx, row_data in enumerate(io_rows):
    row = tbl2.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        cell = row.cells[c_idx]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(9)
        run.font.name = "Garamond"

for row in tbl2.rows:
    row.cells[0].width = Inches(1.5)
    row.cells[1].width = Inches(4.8)

doc.add_paragraph()
body("Outputs (main panel):")

tbl3 = doc.add_table(rows=4, cols=2)
tbl3.style = "Table Grid"
tbl3.alignment = WD_TABLE_ALIGNMENT.LEFT

out_headers = ["Output", "Description"]
out_rows = [
    ["Hourly forecast chart",
     "Bar chart of predicted total passengers by hour (4am–11pm), overlaid with the historical average for the selected airport/month/day-of-week as a dotted reference line"],
    ["Checkpoint breakdown",
     "Horizontal bar chart showing predicted volume per checkpoint at the forecast peak hour"],
    ["Summary metrics",
     "Peak hour, peak predicted volume, total daily forecast, and % deviation from historical average"],
]

for i, h in enumerate(out_headers):
    cell = tbl3.rows[0].cells[i]
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(h)
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.name = "Garamond"
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8E8E8')
    tcPr.append(shd)

for r_idx, row_data in enumerate(out_rows):
    row = tbl3.rows[r_idx + 1]
    for c_idx, val in enumerate(row_data):
        cell = row.cells[c_idx]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(9)
        run.font.name = "Garamond"

for row in tbl3.rows:
    row.cells[0].width = Inches(1.5)
    row.cells[1].width = Inches(4.8)

doc.add_paragraph()

sub_heading("Demo Flow")

body(
    "Demo sequence: select ATL / Friday / June, click Generate Forecast, walk through "
    "the hourly chart (morning peak ~7–8am), show the per-checkpoint breakdown, then "
    "switch to a holiday date to show forecast adjustment. Full flow runs under two "
    "minutes. Screenshots are included in Appendix B (Figures B1 and B2). Note: flight "
    "schedule inputs are entered manually in the POC; a production system would pull "
    "published airline schedule data automatically."
)

# ═══════════════════════════════════════════════════════════════════════════════
# APPENDIX
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()

p = para("", space_before=0, space_after=4)
add_text(p, "A P P E N D I X", size=8, color=GREY)

add_rule()

# Appendix A
p = para("", space_before=10, space_after=4)
add_text(p, "Appendix A: AWS Architecture Diagram", size=11, bold=True)

body(
    "Figure A1. AWS deployment architecture for the TSA Throughput Forecasting POC. "
    "Raw data and model artifact are stored in Amazon S3. An EC2 instance (t3.medium) "
    "hosts the Streamlit application, loads the model from S3 at startup, and serves "
    "live predictions."
)

p = para("", space_before=4, space_after=4)
add_text(p, "[Insert Figure A1: midpoint_visual_05_aws_architecture.png]",
         size=10, italic=True, color=GREY)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Appendix B
p = para("", space_before=14, space_after=4)
add_text(p, "Appendix B: Streamlit Application Screenshots", size=11, bold=True)

body(
    "Figure B1. Main forecast view: hourly throughput forecast (blue bars) vs. historical "
    "average (red dotted line) for a selected airport, month, and day of week."
)

p = para("", space_before=4, space_after=10)
add_text(p,
    "[Insert Figure B1: screenshot of the full Streamlit page (sidebar visible on left, "
    "hourly bar chart visible on right). Select ATL, Friday, June before taking the screenshot.]",
    size=10, italic=True, color=GREY)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

body(
    "Figure B2. Checkpoint breakdown view: predicted passenger volume by checkpoint "
    "at the forecast peak hour."
)

p = para("", space_before=4, space_after=4)
add_text(p,
    "[Insert Figure B2: screenshot scrolled down to show the checkpoint horizontal bar chart "
    "and the three summary metric boxes below it.]",
    size=10, italic=True, color=GREY)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── Works Cited ───────────────────────────────────────────────────────────────
doc.add_paragraph()
add_rule()

p = para("", space_before=6, space_after=2)
add_text(p, "W O R K S  C I T E D", size=8, color=GREY)

p = para("", space_before=4, space_after=0)
add_text(p,
    "Monmousseau, P., Marzuoli, A., Feron, E., & Delahaye, D. (2020). Predicting and "
    "controlling airport passenger flow. ", size=10, italic=False)
add_text(p,
    "Transportation Research Part C: Emerging Technologies, ",
    size=10, italic=True)
add_text(p,
    "120, 102796. https://doi.org/10.1016/j.trc.2020.102796", size=10)

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save("ML_Memo2.docx")
print("Saved: ML_Memo2.docx")
